"""Publish/validate DXY-MID result tables; never run a measurement algorithm.

Inputs are the already-frozen DXY-MID frame/video/group results and confirmed GT.
No S5/G3 result or historical fallback is read. Fresh measurements use dxy-mid.
"""
from pathlib import Path
import argparse
import hashlib
import json

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
LABELS = {
 'group_id':'测量组ID', 'display_group':'更正后显示名称', 'paris':'实际Paris形态',
 'ring':'参考环规格', 'GT_mm':'GT（mm）', 'DXY_MID_mm':'DXY-MID（mm）',
 'signed_error_mm':'有符号误差（mm）', 'absolute_error_mm':'绝对误差（mm）',
 'signed_error_pct':'有符号相对误差（%）', 'absolute_error_pct':'绝对相对误差（%）',
 'n_input_frames':'输入帧数', 'n_valid_frames':'有效帧数', 'n_ring_failed_frames':'Ring失败帧数',
 'n_normal_frames':'Normal帧数', 'n_Q_recovery_frames':'Q恢复帧数', 'n_strict_frames':'Strict帧数',
 'n_input_videos':'输入视频数', 'n_valid_videos':'有效视频数', 'crosses_10mm':'是否错误跨越10mm',
 'classification_10mm':'10mm分类', 'history_prediction_fallback':'历史预测回退',
 'video':'视频ID', 'IQR_mm':'视频IQR（mm）', 'reason':'无效原因',
}


def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def boolean(x):
    text=x.astype(str).str.lower()
    if not text.isin(['true','false','0','1']).all(): raise ValueError('Invalid/missing quality boolean')
    return text.isin(['true','1'])


def calculate(results, gt_path, input_manifest):
    from dxy_mid import _engine
    frames=pd.read_csv(results/'frame_measurements.csv')
    frozen=pd.read_csv(results/'frozen_groups.csv')
    frozen_videos=pd.read_csv(results/'frozen_videos.csv')
    gt=pd.read_csv(gt_path);inputs=pd.read_csv(input_manifest)
    assert len(frames)==790 and frames.source_id.is_unique
    assert set(frames.source_id)==set(inputs.source_id)
    assert len(frozen)==len(gt)==39 and frozen.group.is_unique and gt.group.is_unique
    assert set(frozen.group)==set(gt.group)==set(frames.group)
    assert (frames.history_prediction_fallback==0).all()
    # Independent consistency check against the existing unmodified Fusion.
    regroups,revideos=_engine.fusion(frames.to_dict('records'))
    j=frozen.merge(regroups,on='group',validate='one_to_one',suffixes=('_frozen','_check'))
    np.testing.assert_allclose(j.prediction_mm_frozen,j.prediction_mm_check,rtol=1e-12,atol=1e-10)
    v=frozen_videos.merge(revideos,on=['group','video'],validate='one_to_one',suffixes=('_frozen','_check'))
    assert len(v)==len(frozen_videos)==79
    np.testing.assert_allclose(v.median_mm_frozen,v.median_mm_check,rtol=1e-12,atol=1e-10,equal_nan=True)
    f=frames.copy()
    for column in ['measurement_eligible','ring_valid','normal_Gate_pass','Q_conflict_recovery','strict_stable']:
        f[column]=boolean(f[column])
    f['ring_failed']=~f.ring_valid
    counts=f.groupby('group').agg(n_input_frames=('source_id','size'),n_valid_frames=('measurement_eligible','sum'),
        n_ring_failed_frames=('ring_failed','sum'),n_normal_frames=('normal_Gate_pass','sum'),
        n_Q_recovery_frames=('Q_conflict_recovery','sum'),n_strict_frames=('strict_stable','sum')).reset_index()
    ring=inputs.groupby('group').ring_mm.agg(['first','nunique']);assert ring['nunique'].eq(1).all()
    d=gt[['group','display_group','paris','GT_mm']].merge(frozen,on='group',validate='one_to_one').merge(counts,on='group',validate='one_to_one')
    d['ring']=d.group.map(ring['first']).map({5.:'R5',10.:'R10'});assert d.ring.notna().all()
    d=d.rename(columns={'group':'group_id','prediction_mm':'DXY_MID_mm'})
    d['signed_error_mm']=d.DXY_MID_mm-d.GT_mm
    d['absolute_error_mm']=d.signed_error_mm.abs()
    d['signed_error_pct']=100*d.signed_error_mm/d.GT_mm
    d['absolute_error_pct']=d.signed_error_pct.abs()
    pos=d.GT_mm>=10;pred=d.DXY_MID_mm>=10
    d['crosses_10mm']=pos!=pred
    d['classification_10mm']=np.select([pos&pred,pos&~pred,~pos&pred],['TP','FN','FP'],default='TN')
    columns=list(LABELS)[:21]
    d=d[columns].sort_values(['absolute_error_pct','group_id'],ascending=[False,True]).reset_index(drop=True)
    vv=frozen_videos.merge(gt[['group','paris','GT_mm']],on='group',validate='many_to_one')
    cc=f.groupby(['group','video']).agg(n_ring_failed_frames=('ring_failed','sum'),n_normal_frames=('normal_Gate_pass','sum')).reset_index()
    vv=vv.merge(cc,on=['group','video'],validate='one_to_one').rename(columns={'group':'group_id','median_mm':'DXY_MID_mm','Q_conflict_recovery_frames':'n_Q_recovery_frames','strict_frames':'n_strict_frames'})
    vv['signed_error_mm']=vv.DXY_MID_mm-vv.GT_mm;vv['absolute_error_pct']=100*vv.signed_error_mm.abs()/vv.GT_mm
    metrics=dict(mainline_id='DXY-MID',target_groups=39,valid_groups=int(np.isfinite(d.DXY_MID_mm).sum()),
        input_frames=790,valid_frames=int(f.measurement_eligible.sum()),ring_failed_frames=int(f.ring_failed.sum()),
        normal_frames=int(f.normal_Gate_pass.sum()),Q_recovery_frames=int(f.Q_conflict_recovery.sum()),strict_frames=int(f.strict_stable.sum()),
        median_absolute_relative_error_pct=float(d.absolute_error_pct.median()),MAE_mm=float(d.absolute_error_mm.mean()),
        p95_absolute_relative_error_pct=float(d.absolute_error_pct.quantile(.95)),max_absolute_relative_error_pct=float(d.absolute_error_pct.max()),
        within_10pct_count=int(d.absolute_error_pct.le(10).sum()),within_10pct_percent=float(100*d.absolute_error_pct.le(10).mean()),
        n_videos=79,n_valid_videos=int(np.isfinite(vv.DXY_MID_mm).sum()),threshold10_misclassified=int(d.crosses_10mm.sum()),
        historical_prediction_fallback=0,GT_used_for_algorithm_selection=False)
    assert [metrics[k] for k in ['valid_groups','valid_frames','ring_failed_frames','normal_frames','Q_recovery_frames','strict_frames']]==[39,757,33,627,130,616]
    assert abs(metrics['median_absolute_relative_error_pct']-4.411327)<1e-6 and abs(metrics['MAE_mm']-.709224)<5e-7
    return d,vv,metrics


def workbook(groups, videos, metrics, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.utils import get_column_letter
    wb=Workbook();wb.remove(wb.active)
    for name,data in [('39组误差排序',groups),('逐视频测量',videos),('总体指标',pd.DataFrame([metrics]))]:
        ws=wb.create_sheet(name);ws.append([LABELS.get(c,c) for c in data.columns])
        for row in data.itertuples(index=False,name=None):
            ws.append([None if pd.isna(x) else x.item() if isinstance(x,np.generic) else x for x in row])
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='22577A');cell.alignment=Alignment(wrap_text=True,vertical='center')
        ws.row_dimensions[1].height=38
        for i,c in enumerate(data.columns,1):
            ws.column_dimensions[get_column_letter(i)].width=43 if c in ['group_id','display_group','video'] else 19
            for cell in list(ws.columns)[i-1][1:]:
                if isinstance(cell.value,float):cell.number_format='0.000000'
        if 'absolute_error_pct' in data:
            c=get_column_letter(data.columns.get_loc('absolute_error_pct')+1)
            ws.conditional_formatting.add(f'{c}2:{c}{len(data)+1}',ColorScaleRule(start_type='min',start_color='D8F3DC',end_type='max',end_color='FFADAD'))
    ws=wb.create_sheet('阅读说明')
    for row in [('正式方法','DXY-MID：V9.1 Geometry + Mid Scale'),('有符号误差','prediction - GT；负值低估，正值高估'),('百分比','相对误差百分数已乘100；例如4.41代表4.41%'),('聚合','逐帧毫米→各视频中位数→视频等权组中位数'),('质量','Normal与Q恢复可参与正式聚合；Strict仅作诊断'),('GT','采用人工确认映射；组ID保留冻结关联，实际Paris和显示名已更正'),('研究范围','39个硅胶模型回顾性测量组，非临床泛化性能')]:ws.append(row)
    ws.column_dimensions['A'].width=20;ws.column_dimensions['B'].width=100;ws.freeze_panes='A2'
    wb.save(path)


def verify_files(groups,videos,metrics,output):
    from openpyxl import load_workbook
    for name,expected in [('group_results.csv',groups),('video_results.csv',videos),('metrics_summary.csv',pd.DataFrame([metrics]))]:
        actual=pd.read_csv(output/'tables'/name)
        pd.testing.assert_frame_equal(actual,expected.reset_index(drop=True),check_dtype=False,rtol=1e-11,atol=1e-10)
    wb=load_workbook(output/'tables/group_results.xlsx',read_only=False,data_only=True)
    ws=wb['39组误差排序'];assert ws.freeze_panes=='A2' and ws.max_row==40
    assert [c.value for c in ws[1]]==[LABELS.get(c,c) for c in groups.columns]
    for row,expected in zip(ws.iter_rows(min_row=2,values_only=True),groups.itertuples(index=False,name=None)):
        for a,b in zip(row,expected):
            if isinstance(b,(float,np.floating)):assert np.isclose(a,b,atol=1e-10,rtol=1e-12)
            else:assert a==b,(a,b)
    assert groups.absolute_error_pct.is_monotonic_decreasing
    return dict(passed=True,groups=39,frames=790,CSV_Excel_same_values=True,frozen_predictions_match=True,video_equal_weight_verified=True)


def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--results-root',type=Path,default=ROOT/'validation/results')
    ap.add_argument('--gt',type=Path,default=ROOT/'validation/frozen39/GT_mapping.csv')
    ap.add_argument('--input-manifest',type=Path,default=ROOT/'validation/frozen39/inputs.csv')
    ap.add_argument('--output-root',type=Path,default=ROOT/'mainline')
    ap.add_argument('--check',action='store_true');args=ap.parse_args()
    d,v,m=calculate(args.results_root,args.gt,args.input_manifest)
    if not args.check:
        (args.output_root/'tables').mkdir(parents=True,exist_ok=True)
        d.to_csv(args.output_root/'tables/group_results.csv',index=False)
        v.to_csv(args.output_root/'tables/video_results.csv',index=False)
        pd.DataFrame([m]).to_csv(args.output_root/'tables/metrics_summary.csv',index=False)
        workbook(d,v,m,args.output_root/'tables/group_results.xlsx')
    result=verify_files(d,v,m,args.output_root);print(json.dumps(dict(validation=result,metrics=m),ensure_ascii=False,indent=2))


if __name__=='__main__':main()
